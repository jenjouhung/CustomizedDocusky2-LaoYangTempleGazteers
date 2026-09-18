#!/usr/bin/env python3
"""依 Excel 標記清單，批次標記另一份 Excel 的文字欄位。

標記優先順序：標記清單活頁簿的工作表順序，再依各工作表的資料列順序。
已由較早規則標記的片段不會再被後續規則處理，因此不會產生巢狀標記。
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from xml.sax.saxutils import escape

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    if exc.name != "openpyxl":
        raise
    raise SystemExit(
        "缺少 openpyxl。請先在終端機執行：\n"
        "python3 -m pip install -r requirements-tagger.txt\n"
        "安裝完成後，再執行 python3 tools/text_tagger.py"
    ) from exc


PROJECT_ROOT = Path(__file__).resolve().parent.parent
PROFILE = json.loads((PROJECT_ROOT / "dataset" / "profile.json").read_text(encoding="utf-8"))
TAGGER_DEFAULTS = PROFILE.get("tools", {}).get("textTagger", {})


def _project_path(value: str | None) -> Path | None:
    if value is None:
        return None
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


TARGET_FILE = _project_path(TAGGER_DEFAULTS.get("targetFile"))
TARGET_SHEET: str | None = TAGGER_DEFAULTS.get("targetSheet")
TARGET_COLUMN = TAGGER_DEFAULTS.get("targetColumn") or PROFILE["roles"]["fullText"]
TAG_LIST_FILE = _project_path(TAGGER_DEFAULTS.get("tagListFile"))
OUTPUT_FILE = _project_path(TAGGER_DEFAULTS.get("outputFile"))

ALLOWED_HEADERS = {"tagname", "tagval", "@term", "@refid"}
XML_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_.-]*$")
EXISTING_TAG_RE = re.compile(r"</?[A-Za-z_][A-Za-z0-9_.-]*(?:\s[^<>]*)?>")


@dataclass(frozen=True)
class TagRule:
    tag_name: str
    tag_value: str
    term: str | None
    ref_id: str | None
    sheet_name: str
    row_number: int

    def opening_tag(self) -> str:
        def attribute(value: str) -> str:
            return f'"{escape(value, {chr(34): "&quot;"})}"'

        attributes: list[str] = []
        if self.term is not None:
            attributes.append(f"term={attribute(self.term)}")
        if self.ref_id is not None:
            attributes.append(f"RefId={attribute(self.ref_id)}")
        suffix = f" {' '.join(attributes)}" if attributes else ""
        return f"<{self.tag_name}{suffix}>"

    def wrap(self, text: str) -> str:
        return f"{self.opening_tag()}{text}</{self.tag_name}>"


@dataclass
class TaggingSummary:
    source_rows: int = 0
    nonempty_cells: int = 0
    changed_cells: int = 0
    total_tags: int = 0


def _cell_text(value: object) -> str | None:
    """將規則儲存格轉成字串；空白視為 None。"""
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def _canonical_headers(values: Iterable[object]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for index, value in enumerate(values):
        text = _cell_text(value)
        if text is None:
            continue
        canonical = text.casefold()
        if canonical not in ALLOWED_HEADERS:
            raise ValueError(f"不支援的標記清單欄位：{text!r}")
        if canonical in headers:
            raise ValueError(f"標記清單欄位重複：{text!r}")
        headers[canonical] = index
    return headers


def load_tag_rules(path: Path) -> list[TagRule]:
    """依工作表順序及列順序載入規則。工作表名稱不參與判斷。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    rules: list[TagRule] = []

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue

            try:
                headers = _canonical_headers(header_row)
            except ValueError as exc:
                raise ValueError(f"工作表 {worksheet.title!r}：{exc}") from exc

            missing = {"tagname", "tagval"} - headers.keys()
            if missing:
                names = ", ".join(sorted(missing))
                raise ValueError(
                    f"工作表 {worksheet.title!r} 缺少必要欄位：{names}"
                )

            for row_number, row in enumerate(rows, start=2):
                def get(header: str) -> str | None:
                    index = headers.get(header)
                    return _cell_text(row[index]) if index is not None and index < len(row) else None

                tag_name = get("tagname")
                tag_value = get("tagval")
                term = get("@term") or tag_value
                ref_id = get("@refid")

                if tag_name is None and tag_value is None and term is None and ref_id is None:
                    continue
                if tag_name is None or tag_value is None:
                    raise ValueError(
                        f"工作表 {worksheet.title!r} 第 {row_number} 列："
                        "tagName 與 tagVal 不得空白"
                    )
                if not XML_NAME_RE.fullmatch(tag_name):
                    raise ValueError(
                        f"工作表 {worksheet.title!r} 第 {row_number} 列："
                        f"tagName {tag_name!r} 不是有效的 XML 元素名稱"
                    )

                rules.append(
                    TagRule(
                        tag_name=tag_name,
                        tag_value=tag_value,
                        term=term,
                        ref_id=ref_id,
                        sheet_name=worksheet.title,
                        row_number=row_number,
                    )
                )
    finally:
        workbook.close()

    if not rules:
        raise ValueError("標記清單中沒有可用規則")
    return rules


def tag_text(text: str, rules: Iterable[TagRule]) -> tuple[str, int]:
    """標記純文字；較早規則產生的片段不再接受後續規則。"""
    # (內容, 是否已標記)。已標記片段包含完整元素字串，後續不再搜尋。
    segments: list[tuple[str, bool]] = [(text, False)]
    tag_count = 0

    for rule in rules:
        updated: list[tuple[str, bool]] = []
        for content, protected in segments:
            if protected or rule.tag_value not in content:
                updated.append((content, protected))
                continue

            parts = content.split(rule.tag_value)
            for index, part in enumerate(parts):
                if part:
                    updated.append((part, False))
                if index < len(parts) - 1:
                    updated.append((rule.wrap(rule.tag_value), True))
                    tag_count += 1
        segments = updated

    return "".join(content for content, _ in segments), tag_count


def _find_column(worksheet, column_name: str) -> int:
    matches = [
        cell.column
        for cell in worksheet[1]
        if _cell_text(cell.value) == column_name.strip()
    ]
    if not matches:
        raise ValueError(
            f"工作表 {worksheet.title!r} 的第 1 列找不到欄位 {column_name!r}"
        )
    if len(matches) > 1:
        raise ValueError(
            f"工作表 {worksheet.title!r} 的欄位 {column_name!r} 重複出現"
        )
    return matches[0]


def tag_workbook(
    target_file: Path,
    tag_list_file: Path,
    output_file: Path,
    target_column: str,
    target_sheet: str | None = None,
) -> TaggingSummary:
    target_file = target_file.expanduser().resolve()
    tag_list_file = tag_list_file.expanduser().resolve()
    output_file = output_file.expanduser().resolve()

    if target_file == output_file:
        raise ValueError("輸出檔不得覆寫來源檔，請指定不同的 OUTPUT_FILE")
    if not target_file.is_file():
        raise FileNotFoundError(f"找不到目標 Excel：{target_file}")
    if not tag_list_file.is_file():
        raise FileNotFoundError(f"找不到標記清單 Excel：{tag_list_file}")

    rules = load_tag_rules(tag_list_file)
    workbook = load_workbook(target_file)
    summary = TaggingSummary()

    try:
        if target_sheet is None:
            worksheet = workbook.worksheets[0]
        elif target_sheet in workbook.sheetnames:
            worksheet = workbook[target_sheet]
        else:
            raise ValueError(f"目標 Excel 找不到工作表 {target_sheet!r}")

        column_index = _find_column(worksheet, target_column)
        summary.source_rows = max(worksheet.max_row - 1, 0)

        for row_number in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_number, column=column_index)
            if cell.value is None or str(cell.value) == "":
                continue
            if not isinstance(cell.value, str):
                raise ValueError(
                    f"工作表 {worksheet.title!r} 第 {row_number} 列的 "
                    f"{target_column!r} 不是文字"
                )
            if EXISTING_TAG_RE.search(cell.value):
                raise ValueError(
                    f"工作表 {worksheet.title!r} 第 {row_number} 列似乎已有標記；"
                    "為避免重複或巢狀標記，已停止處理"
                )

            summary.nonempty_cells += 1
            tagged, count = tag_text(cell.value, rules)
            if count:
                cell.value = tagged
                summary.changed_cells += 1
                summary.total_tags += count

        output_file.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_file)
    finally:
        workbook.close()

    return summary


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target-file", type=Path, default=TARGET_FILE, required=TARGET_FILE is None)
    parser.add_argument("--target-sheet", default=TARGET_SHEET)
    parser.add_argument("--target-column", default=TARGET_COLUMN)
    parser.add_argument("--tag-list-file", type=Path, default=TAG_LIST_FILE, required=TAG_LIST_FILE is None)
    parser.add_argument("--output-file", type=Path, default=OUTPUT_FILE, required=OUTPUT_FILE is None)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    summary = tag_workbook(
        target_file=args.target_file,
        tag_list_file=args.tag_list_file,
        output_file=args.output_file,
        target_column=args.target_column,
        target_sheet=args.target_sheet,
    )
    print(f"處理完成：{args.output_file.resolve()}")
    print(f"資料列數：{summary.source_rows}")
    print(f"非空白文字：{summary.nonempty_cells}")
    print(f"已修改儲存格：{summary.changed_cells}")
    print(f"新增標記總數：{summary.total_tags}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
