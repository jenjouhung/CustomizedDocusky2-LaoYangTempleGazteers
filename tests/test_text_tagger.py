import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools.text_tagger import TagRule, load_tag_rules, tag_text, tag_workbook


def rule(name, value, term=None, ref_id=None, row=2):
    return TagRule(name, value, term, ref_id, "rules", row)


class TagTextTests(unittest.TestCase):
    def test_example_and_attributes(self):
        source = "中明寺比丘尼道揚、道積、道保依方等行道"
        actual, count = tag_text(
            source, [rule("person", "道揚", "揚法師", "A000123")]
        )
        self.assertEqual(
            actual,
            '中明寺比丘尼<person term="揚法師" RefId="A000123">道揚</person>、道積、道保依方等行道',
        )
        self.assertEqual(count, 1)

    def test_earlier_rule_wins_and_no_nested_markup(self):
        actual, count = tag_text(
            "新城縣功曹",
            [rule("place", "新城縣"), rule("office", "縣功曹")],
        )
        self.assertEqual(actual, "<place>新城縣</place>功曹")
        self.assertEqual(count, 1)

    def test_all_repeated_occurrences_are_tagged(self):
        actual, count = tag_text("彌勒與彌勒", [rule("term", "彌勒")])
        self.assertEqual(actual, "<term>彌勒</term>與<term>彌勒</term>")
        self.assertEqual(count, 2)

    def test_attribute_values_are_escaped(self):
        actual, _ = tag_text("甲", [rule("person", "甲", 'A&B"C')])
        self.assertEqual(actual, '<person term="A&amp;B&quot;C">甲</person>')


class WorkbookTests(unittest.TestCase):
    def test_rule_order_across_sheets_and_case_insensitive_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "rules.xlsx"
            workbook = Workbook()
            first = workbook.active
            first.title = "任意名稱一"
            first.append(["tagName", "tagVal", "@Term"])
            first.append(["person", "道揚", "揚法師"])
            second = workbook.create_sheet("任意名稱二")
            second.append(["tagName", "tagVal", "@refid"])
            second.append(["place", "中明寺", "P1"])
            workbook.save(path)

            rules = load_tag_rules(path)
            self.assertEqual([item.tag_value for item in rules], ["道揚", "中明寺"])
            self.assertEqual(rules[0].term, "揚法師")
            self.assertEqual(rules[1].term, "中明寺")
            self.assertEqual(rules[1].ref_id, "P1")

    def test_workbook_output_preserves_other_cells(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            target = directory / "target.xlsx"
            rules_path = directory / "rules.xlsx"
            output = directory / "output.xlsx"

            target_book = Workbook()
            target_sheet = target_book.active
            target_sheet.append(["編號", "全文"])
            target_sheet.append([1, "道揚與道積"])
            target_book.save(target)

            rules_book = Workbook()
            rule_sheet = rules_book.active
            rule_sheet.append(["tagName", "tagVal"])
            rule_sheet.append(["person", "道揚"])
            rules_book.save(rules_path)

            summary = tag_workbook(target, rules_path, output, "全文")
            result = load_workbook(output, read_only=True)
            self.assertEqual(result.active["A2"].value, 1)
            self.assertEqual(
                result.active["B2"].value,
                '<person term="道揚">道揚</person>與道積',
            )
            self.assertEqual(summary.changed_cells, 1)
            result.close()


if __name__ == "__main__":
    unittest.main()
